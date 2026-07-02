import os
import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest import mock

from document_loader import (
    MDLoader,
    PDFLoader,
    TXTLoader,
    WordLoader,
    get_document_loader,
    get_supported_extensions,
    load_and_split_document,
    load_and_split_documents,
    load_document,
    load_documents,
)


FIXTURES_DIR = Path(__file__).parent / "fixtures"


class DocumentLoaderEntrypointTests(unittest.TestCase):
    def test_supported_extensions_include_core_formats(self):
        extensions = get_supported_extensions()

        self.assertIn(".txt", extensions)
        self.assertIn(".md", extensions)
        self.assertIn(".markdown", extensions)
        self.assertIn(".pdf", extensions)
        self.assertIn(".docx", extensions)

    def test_get_document_loader_selects_loader_by_extension(self):
        self.assertIsInstance(get_document_loader(str(FIXTURES_DIR / "sample.txt")), TXTLoader)
        self.assertIsInstance(get_document_loader(str(FIXTURES_DIR / "sample.md")), MDLoader)

    def test_get_document_loader_ignores_loader_specific_kwargs_for_other_formats(self):
        loader = get_document_loader(
            str(FIXTURES_DIR / "sample.txt"),
            extract_tables=True,
            ocr_enabled=True,
            ocr_min_chars=1,
            ocr_dpi=200,
        )

        self.assertIsInstance(loader, TXTLoader)

    def test_get_document_loader_rejects_unsupported_extension(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.csv"
            path.write_text("a,b,c", encoding="utf-8")

            with self.assertRaises(ValueError):
                get_document_loader(str(path))

    def test_load_document_loads_txt_with_metadata(self):
        document = load_document(str(FIXTURES_DIR / "sample.txt"), clean=True)

        self.assertIn("RAG ingestion prototype", document.content)
        self.assertEqual(document.metadata["file_type"], "txt")
        self.assertTrue(document.metadata["cleaned"])
        self.assertIn("source", document.metadata)

    def test_load_document_loads_markdown_front_matter(self):
        document = load_document(str(FIXTURES_DIR / "sample.md"), clean=True)

        self.assertEqual(document.metadata["file_type"], "markdown")
        self.assertEqual(document.metadata["title"], "Sample Markdown Fixture")
        self.assertTrue(document.metadata["has_front_matter"])
        self.assertEqual(document.metadata["code_block_count"], 1)
        self.assertIn("Knowledge Base Notes", document.content)
        self.assertIn("project notes", document.content)

    def test_load_documents_loads_supported_files_from_directory(self):
        documents = load_documents(str(FIXTURES_DIR), clean=True)
        file_types = sorted(document.metadata["file_type"] for document in documents)

        self.assertEqual(file_types, ["markdown", "txt"])

    def test_load_and_split_document_adds_chunk_metadata(self):
        chunks = load_and_split_document(
            str(FIXTURES_DIR / "sample.txt"),
            clean=True,
            chunk_size=80,
            chunk_overlap=10,
        )

        self.assertGreaterEqual(len(chunks), 2)
        self.assertEqual(chunks[0].metadata["chunk_index"], 0)
        self.assertEqual(chunks[0].metadata["total_chunks"], len(chunks))
        self.assertIn("start_char", chunks[0].metadata)
        self.assertIn("end_char", chunks[0].metadata)
        self.assertEqual(chunks[0].metadata["loader_entrypoint"], "load_and_split_document")

    def test_load_and_split_documents_splits_directory(self):
        chunks = load_and_split_documents(
            str(FIXTURES_DIR),
            clean=True,
            chunk_size=90,
            chunk_overlap=10,
        )

        sources = {Path(chunk.metadata["source"]).name for chunk in chunks}
        self.assertEqual(sources, {"sample.md", "sample.txt"})
        self.assertTrue(all("chunk_index" in chunk.metadata for chunk in chunks))


class GeneratedDocumentLoaderTests(unittest.TestCase):
    def _write_text_pdf(self, path: Path, page_texts):
        import fitz

        doc = fitz.open()
        for text in page_texts:
            page = doc.new_page()
            if text:
                page.insert_text((72, 72), text, fontsize=14)
        doc.save(path)
        doc.close()

    def _make_pdf_text_image(self, text: str) -> bytes:
        import fitz

        image_doc = fitz.open()
        page = image_doc.new_page(width=480, height=160)
        page.insert_text((40, 90), text, fontsize=28)
        pixmap = page.get_pixmap(alpha=False)
        image_bytes = pixmap.tobytes("png")
        image_doc.close()
        return image_bytes

    def _write_scanned_pdf(self, path: Path, image_text: str):
        import fitz

        doc = fitz.open()
        page = doc.new_page(width=480, height=160)
        page.insert_image(page.rect, stream=self._make_pdf_text_image(image_text))
        doc.save(path)
        doc.close()

    def _write_mixed_pdf(self, path: Path, text_page: str, scanned_text: str):
        import fitz

        doc = fitz.open()
        page = doc.new_page(width=480, height=160)
        page.insert_text((40, 90), text_page, fontsize=18)
        scanned_page = doc.new_page(width=480, height=160)
        scanned_page.insert_image(scanned_page.rect, stream=self._make_pdf_text_image(scanned_text))
        doc.save(path)
        doc.close()

    def test_word_loader_with_generated_docx(self):
        try:
            from docx import Document as DocxDocument
        except ImportError:
            self.skipTest("python-docx is not installed")

        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.docx"
            doc = DocxDocument()
            doc.add_paragraph("Generated Word fixture for loader verification.")
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Key"
            table.cell(0, 1).text = "Value"
            doc.save(path)

            loader = get_document_loader(str(path))
            loaded = loader.load_with_metadata()

            self.assertIsInstance(loader, WordLoader)
            self.assertEqual(loaded.metadata["file_type"], "docx")
            self.assertIn("Generated Word fixture", loaded.content)
            self.assertIn("[Table 1]", loaded.content)

    def test_pdf_loader_with_generated_pdf(self):
        try:
            import fitz
        except ImportError:
            self.skipTest("PyMuPDF is not installed")

        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.pdf"
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((72, 72), "Generated PDF fixture for loader verification.")
            doc.save(path)
            doc.close()

            loader = get_document_loader(str(path))
            loaded = loader.load_with_metadata()

            self.assertIsInstance(loader, PDFLoader)
            self.assertEqual(loaded.metadata["file_type"], "pdf")
            self.assertEqual(loaded.metadata["total_pages"], 1)
            self.assertIn("Generated PDF fixture", loaded.content)

    def test_pdf_loader_detects_scanned_page_without_ocr(self):
        try:
            import fitz
        except ImportError:
            self.skipTest("PyMuPDF is not installed")

        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "scan.pdf"
            self._write_scanned_pdf(path, "SCANNEDTOKEN001")

            doc = fitz.open(path)
            try:
                self.assertEqual(doc[0].get_text().strip(), "")
            finally:
                doc.close()

            loader = PDFLoader(str(path), ocr_enabled=False)
            with self.assertLogs("document_loader.pdf_loader", level="WARNING") as logs:
                loaded = loader.load_with_metadata()

            self.assertEqual(loaded.content, "--- Page 1 ---")
            self.assertEqual(loaded.metadata["scanned_page_count"], 1)
            self.assertEqual(loaded.metadata["ocr_page_count"], 0)
            self.assertIn("PDF scanned pages detected", "\n".join(logs.output))

    def test_pdf_loader_ocr_is_per_page_for_mixed_pdf(self):
        try:
            import fitz  # noqa: F401
        except ImportError:
            self.skipTest("PyMuPDF is not installed")

        class SpyOCRPDFLoader(PDFLoader):
            def __init__(self, *args, **kwargs):
                super().__init__(*args, **kwargs)
                self.ocr_calls = 0

            def _ocr_page(self, page):
                self.ocr_calls += 1
                return "OCRUNIQUEPAGE002"

        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "mixed.pdf"
            self._write_mixed_pdf(path, "Embedded text page.", "OCRUNIQUEPAGE002")

            loader = SpyOCRPDFLoader(str(path), ocr_enabled=True)
            loaded = loader.load_with_metadata()

            self.assertIn("Embedded text page.", loaded.content)
            self.assertIn("OCRUNIQUEPAGE002", loaded.content)
            self.assertEqual(loader.ocr_calls, 1)
            self.assertEqual(loaded.metadata["scanned_page_count"], 1)
            self.assertEqual(loaded.metadata["ocr_page_count"], 1)

    def test_pdf_loader_raises_when_ocr_enabled_and_dependency_missing(self):
        try:
            import fitz  # noqa: F401
        except ImportError:
            self.skipTest("PyMuPDF is not installed")

        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "scan.pdf"
            self._write_scanned_pdf(path, "SCANNEDTOKEN002")
            loader = PDFLoader(str(path), ocr_enabled=True)

            with mock.patch.object(loader, "_get_ocr_engine", side_effect=ImportError("rapidocr missing")):
                with self.assertRaisesRegex(ImportError, "rapidocr missing"):
                    loader.load()

    def test_pdf_loader_table_extraction_opens_pdfplumber_once(self):
        try:
            import fitz  # noqa: F401
        except ImportError:
            self.skipTest("PyMuPDF is not installed")

        class FakePage:
            def __init__(self, tables):
                self._tables = tables

            def extract_tables(self):
                return self._tables

        class FakePlumberPDF:
            def __init__(self):
                self.pages = [
                    FakePage([[["H", "V"], ["A", "B"]]]),
                    FakePage([]),
                    FakePage([[["X"], ["Y"]]]),
                ]
                self.closed = False

            def close(self):
                self.closed = True

        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "tables.pdf"
            self._write_text_pdf(path, ["PDF page 1 text.", "PDF page 2 text.", "PDF page 3 text."])
            fake_pdf = FakePlumberPDF()
            open_calls = []

            def fake_open(open_path):
                open_calls.append(open_path)
                return fake_pdf

            with mock.patch.dict(sys.modules, {"pdfplumber": types.SimpleNamespace(open=fake_open)}):
                loader = PDFLoader(str(path), extract_tables=True)
                content = loader.load()

            expected = "\n".join(
                [
                    "\n--- Page 1 ---\nPDF page 1 text.",
                    "\n[Table]\n| H | V |\n| --- | --- |\n| A | B |\n",
                    "\n--- Page 2 ---\nPDF page 2 text.",
                    "\n--- Page 3 ---\nPDF page 3 text.",
                    "\n[Table]\n| X |\n| --- |\n| Y |\n",
                ]
            ).strip()
            self.assertEqual(content, expected)
            self.assertEqual(open_calls, [path])
            self.assertTrue(fake_pdf.closed)

    def test_pdf_loader_does_not_open_pdfplumber_when_table_extraction_disabled(self):
        try:
            import fitz  # noqa: F401
        except ImportError:
            self.skipTest("PyMuPDF is not installed")

        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "no_tables.pdf"
            self._write_text_pdf(path, ["PDF page 1 text.", "PDF page 2 text."])
            open_calls = []

            def fake_open(open_path):
                open_calls.append(open_path)
                raise AssertionError("pdfplumber should not be opened")

            with mock.patch.dict(sys.modules, {"pdfplumber": types.SimpleNamespace(open=fake_open)}):
                loader = PDFLoader(str(path), extract_tables=False)
                content = loader.load()

            expected = "\n".join(
                [
                    "\n--- Page 1 ---\nPDF page 1 text.",
                    "\n--- Page 2 ---\nPDF page 2 text.",
                ]
            ).strip()
            self.assertEqual(content, expected)
            self.assertEqual(open_calls, [])

    def test_pdf_loader_skips_tables_when_pdfplumber_missing(self):
        try:
            import fitz  # noqa: F401
        except ImportError:
            self.skipTest("PyMuPDF is not installed")

        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "missing_pdfplumber.pdf"
            self._write_text_pdf(path, ["PDF page text."])

            with mock.patch.dict(sys.modules, {"pdfplumber": None}):
                loader = PDFLoader(str(path), extract_tables=True)
                with self.assertLogs("document_loader.pdf_loader", level="WARNING") as logs:
                    content = loader.load()

            self.assertEqual(content, "--- Page 1 ---\nPDF page text.")
            self.assertIn("pdfplumber is not installed", "\n".join(logs.output))

    def test_rapidocr_integration_reads_scanned_pdf_when_enabled(self):
        if os.getenv("RUN_RAPIDOCR_TEST") != "1":
            self.skipTest("Set RUN_RAPIDOCR_TEST=1 to run the gated RapidOCR integration test")
        try:
            import fitz  # noqa: F401
            import rapidocr_onnxruntime  # noqa: F401
        except ImportError:
            self.skipTest("RapidOCR integration dependencies are not installed")

        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "rapidocr.pdf"
            self._write_scanned_pdf(path, "RAPIDOCRTOKEN123")

            loader = PDFLoader(str(path), ocr_enabled=True, ocr_dpi=200)
            loaded = loader.load_with_metadata()

            self.assertIn("RAPIDOCRTOKEN", loaded.content.replace(" ", ""))
            self.assertEqual(loaded.metadata["scanned_page_count"], 1)
            self.assertEqual(loaded.metadata["ocr_page_count"], 1)

    def test_txt_loader_rejects_binary_payload(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "binary.txt"
            path.write_bytes(b"\x00\x01\x02\x03not text")

            with self.assertRaisesRegex(ValueError, "binary"):
                TXTLoader(str(path)).load()

    def test_txt_loader_accepts_gbk_chinese_text(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "gbk.txt"
            expected = "中文知识库"
            path.write_bytes(expected.encode("gbk"))

            self.assertEqual(TXTLoader(str(path)).load(), expected)

    def test_empty_text_document_splits_to_zero_chunks(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "empty.txt"
            path.write_text("", encoding="utf-8")

            chunks = load_and_split_document(str(path), clean=False)

            self.assertEqual(chunks, [])

    def test_zero_page_pdf_splits_to_zero_chunks(self):
        try:
            import fitz  # noqa: F401
        except ImportError:
            self.skipTest("PyMuPDF is not installed")

        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "empty.pdf"
            path.write_bytes(
                b"%PDF-1.4\n"
                b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
                b"2 0 obj<</Type/Pages/Count 0/Kids[]>>endobj\n"
                b"trailer<</Root 1 0 R>>\n"
                b"%%EOF"
            )

            self.assertEqual(PDFLoader(str(path)).load(), "")
            self.assertEqual(load_and_split_document(str(path), clean=False), [])

    def test_empty_word_document_splits_to_zero_chunks(self):
        try:
            from docx import Document as DocxDocument
        except ImportError:
            self.skipTest("python-docx is not installed")

        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "empty.docx"
            doc = DocxDocument()
            doc.save(path)

            self.assertEqual(WordLoader(str(path)).load(), "")
            self.assertEqual(load_and_split_document(str(path), clean=False), [])

    def test_word_loader_preserves_multiple_tables(self):
        try:
            from docx import Document as DocxDocument
        except ImportError:
            self.skipTest("python-docx is not installed")

        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "multi-table.docx"
            doc = DocxDocument()
            doc.add_paragraph("Document with two tables.")
            first = doc.add_table(rows=1, cols=2)
            first.cell(0, 0).text = "First"
            first.cell(0, 1).text = "Table"
            second = doc.add_table(rows=1, cols=2)
            second.cell(0, 0).text = "Second"
            second.cell(0, 1).text = "Table"
            doc.save(path)

            loaded = WordLoader(str(path)).load_with_metadata()

            self.assertIn("[Table 1]", loaded.content)
            self.assertIn("[Table 2]", loaded.content)
            self.assertIn("First | Table", loaded.content)
            self.assertIn("Second | Table", loaded.content)


if __name__ == "__main__":
    unittest.main()
