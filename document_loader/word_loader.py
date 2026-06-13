from typing import Dict

from .base import DocumentLoader
from logger import setup_logger

logger = setup_logger(__name__)


class WordLoader(DocumentLoader):
    """Word document loader for .docx files."""

    def load(self) -> str:
        """Load text content from a Word document."""
        logger.info("Loading Word file | path=%s", self.file_path)

        try:
            from docx import Document as DocxDocument

            doc = DocxDocument(self.file_path)
            texts = []

            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    texts.append(paragraph.text)

            for index, table in enumerate(doc.tables):
                texts.append(f"\n[Table {index + 1}]")
                for row in table.rows:
                    row_text = " | ".join(cell.text for cell in row.cells)
                    texts.append(row_text)

            full_text = "\n\n".join(texts)
            logger.info(
                "Word loaded | paragraphs=%s | chars=%s",
                len(doc.paragraphs),
                len(full_text),
            )
            return full_text

        except ImportError:
            logger.error("python-docx is not installed. Run: pip install python-docx")
            raise

        except Exception as e:
            logger.error("Word loading failed | error=%s", e)
            raise

    def get_metadata(self) -> Dict:
        """Return Word document metadata."""
        from docx import Document as DocxDocument

        doc = DocxDocument(self.file_path)
        return {
            "file_type": "docx",
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
            "file_size": self.file_path.stat().st_size,
        }
